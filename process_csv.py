# -*- coding: utf-8 -*-
"""정보나루 국외도서 CSV -> 영어원서 큐레이션 엑셀 생성."""
import csv, os, sys
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

CSV_PATH = sys.argv[1] if len(sys.argv) > 1 else \
    r"C:\Users\user1\Downloads\BestLoanList_20260921103300\20260921103300\인기대출도서_2026.csv"
OUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "output")
OUT_XLSX = os.path.join(OUT_DIR, "조이네_영어원서_큐레이션_2026.xlsx")

ENGLISH_PREFIXES = ("9780", "9781")

def load():
    with open(CSV_PATH, encoding="utf-8-sig") as f:
        lines = f.readlines()
    start = next(i for i, l in enumerate(lines) if l.startswith("순위"))
    rows = list(csv.reader(lines[start:]))
    header = [h.strip() for h in rows[0]]
    idx = {name: header.index(name) for name in
           ["순위", "서명", "저자", "출판사", "출판년도", "ISBN", "KDC", "대출건수"]}
    data = []
    for r in rows[1:]:
        if len(r) < 8 or not r[0].strip():
            continue
        isbn = r[idx["ISBN"]].replace("-", "").strip()
        if not isbn.startswith(ENGLISH_PREFIXES):
            continue
        try:
            loans = int(r[idx["대출건수"]] or 0)
        except ValueError:
            loans = 0
        data.append({
            "서명": r[idx["서명"]].strip(),
            "저자": r[idx["저자"]].strip(),
            "출판사": r[idx["출판사"]].strip(),
            "출판년도": r[idx["출판년도"]].strip(),
            "ISBN": isbn,
            "대출건수": loans,
        })
    data.sort(key=lambda d: d["대출건수"], reverse=True)
    return data

def build_xlsx(data):
    os.makedirs(OUT_DIR, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "영어원서_대출순위"

    head_fill = PatternFill("solid", fgColor="2F5496")
    head_font = Font(color="FFFFFF", bold=True)
    cols = ["순위", "서명", "저자", "출판사", "출판년도", "ISBN", "대출건수"]
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(1, c)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for rank, d in enumerate(data, 1):
        ws.append([rank, d["서명"], d["저자"], d["출판사"],
                   d["출판년도"], d["ISBN"], d["대출건수"]])

    widths = [6, 45, 30, 28, 8, 16, 8]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{len(data)+1}"

    # 요약 시트
    ws2 = wb.create_sheet("요약")
    ws2.append(["항목", "값"])
    ws2.append(["영어원서 총 종수", len(data)])
    ws2.append(["대출건수 합계", sum(d["대출건수"] for d in data)])
    ws2.append([])
    ws2.append(["출판사 TOP 15 (종수 기준)", ""])
    for pub, cnt in Counter(d["출판사"] for d in data).most_common(15):
        ws2.append([pub, cnt])
    for c in range(1, 3):
        ws2.cell(1, c).font = Font(bold=True)
    ws2.column_dimensions["A"].width = 40
    ws2.column_dimensions["B"].width = 14

    wb.save(OUT_XLSX)
    return OUT_XLSX

if __name__ == "__main__":
    data = load()
    path = build_xlsx(data)
    print(f"영어원서 {len(data)}종 -> {path}")
    print("\n=== 대출 TOP 20 ===")
    for i, d in enumerate(data[:20], 1):
        print(f"{i:2}. [{d['대출건수']:>4}] {d['서명'][:45]}")
